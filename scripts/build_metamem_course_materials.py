"""Build the editable personal reading report and shared slide content from fixed evidence."""
import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "tmp/metamem-materials-20261019"
STEM = "metamem-reading-and-pilot-20261019"
PAPER = "https://aclanthology.org/2026.findings-acl.270/"


def p(text):
    return {"type": "paragraph", "text": text}


def t(headers, rows, widths=None):
    return {"type": "table", "headers": headers, "rows": rows, "widths": widths}


PAGES = [
    ("MetaMem 论文阅读与个人试点阶段报告", [
        p("天津大学 2026 年 NLP 课程设计  第 8 队"),
        p("专题 4 长上下文场景结构化记忆建模  论文组 2"),
        p("汇报人 __________  学号 __________  整理日期 2026 年 10 月 19 日"),
        p("摘要  MetaMem 将可检索的事实记忆 M 与指导事实使用的经验 E 分开建模，通过回答采样、答案反馈、反思和符号编辑积累自然语言规则。其核心问题是正确证据已经出现时，模型仍可能错误处理事件包含关系、时间或数量。原论文在两个大模型上取得平均指标提升；本文同时整理仓库中杨垚负责的 BM25 RAG 个人试点，说明缩小实现的真实结果及限制。"),
        p("现有试点完成规则学习、采样稳定性、检索诊断及回答提示词对照。十道开发题上的结果没有支持稳定收益，角色预算和证据约束提示词均未替换默认设置。引用更多与答案更短也不能替代依据正确。下一步应先核对评分边界，再冻结独立评估方案。"),
        p("关键词  元记忆  长期对话  事实检索  符号优化  依据质量"),
        p("阅读范围  论文问题、方法与原始实验见第 1 至 3 节；仓库已有个人探索见第 4 至 6 节。原始试点负责人为杨垚，本文的整理与汇报署名需本人核对填写。当前材料不代表全组已确定最终方案。"),
    ]),
    ("1 研究问题与基本概念", [
        p("长期对话系统通常把历史交互分块、压缩或组织为可检索记忆，以突破上下文长度限制。但相关片段可能散落在不同会话，局部事实的正确性并不能保证最后推断正确。MetaMem 关注检索后如何使用证据，而不是只增加存储量或检索数量。[1，PDF 第 1 至 3 页]"),
        t(["对象", "含义", "旅行示例"], [["q", "当前问题", "四次旅行共多少里程"], ["M", "与问题相关的事实记忆", "三次共 1800 英里，另一次共 1200 英里，其中首日 300 英里"], ["E", "可迁移的事实使用经验", "判断总量与部分的包含关系，避免重复计数"]], [1.7, 5.5, 9.3]),
        p("论文 Figure 1 中，300 英里已经包含在 1200 英里的总量内，正确计算为 1800 + 1200 = 3000。把所有出现的数值相加得到 3300，错误来自事实关系的使用，不能仅靠列出更多数字修正。[1，PDF 第 1 页 Figure 1]"),
        p("普通记忆回答可写为 y = LLM(prompt(q, M))；MetaMem 增加学习到的经验后为 y = LLM(prompt(q, M, E))。E 保存解决问题的通用策略，M 保存具体事实，两者的作用不同。[1，式 1 至 3，PDF 第 3 页]"),
        p("训练集合包含问题、答案和事实。答案用于训练反馈与独立评分；测试回答只接收问题、事实及冻结经验。验证用于选择规则版本，测试用于评价已固定的方法。主方法更新自然语言规则，没有通过梯度更新回答模型权重，但仍有多次模型调用成本。"),
    ]),
    ("2 元记忆的学习与演化", [
        p("MetaMem 从空经验集合开始，按批处理训练题。每题在相同事实和当前经验下采样多个回答，由更强 judge 与标准答案比较，再反思成功和失败的原因。反思进一步产生规则编辑建议，同批建议经过冲突过滤后更新经验集合。[1，Figure 2，式 4 至 10，PDF 第 3 至 5 页]"),
        {"type": "figure", "path": "metamem-figure2.png", "caption": "图 1  MetaMem 原论文 Figure 2  方法整体流程  来源 [1] PDF 第 4 页"},
        t(["编辑", "作用"], [["ADD", "增加一条经验"], ["DEL", "删除失效或错误经验"], ["MOD", "修正现有经验的适用方式"]], [2.0, 14.5]),
        p("官方固定版本在每题多次采样时，默认只让同一题中既有正确又有错误的组参与更新。全对、全错或评分不完整的组跳过。因此规则没有变化可能是筛选结果；不能悄悄手填规则来替代学习结果。[2，train_metamem.py；3]"),
    ]),
    ("3 原论文的实验与证据", [
        p("论文使用 LongMemEval 的 500 道题，覆盖六类任务。附录规定五折，每折 350 训练、50 验证、100 测试，seed 为 42。正文的 400 道构建题在附录进一步拆为训练与验证。主要事实底座为 LightMem，检索使用 all-MiniLM-L6-v2，top-k 为 20。[1，PDF 第 5、11 页]"),
        p("回答模型为 Qwen3-30B-A3B-Instruct 和 Llama3.1-70B-Instruct，judge 为 Qwen3-235B-A22B。学习采用 5 epochs、batch 50、每题 5 次采样；学习温度 0.7，测试温度 0。它们与本地量化 7B 试点不是相同实验条件。"),
        t(["生成模型", "LightMem Avg", "MetaMem Avg", "差值"], [["Qwen3 30B A3B", "67.50", "71.90", "+4.40 个百分点"], ["Llama3.1 70B", "66.17", "69.08", "+2.91 个百分点"]], [6.0, 3.2, 3.2, 4.1]),
        p("表 1  原论文 Table 1 的 Avg 原值，单位为百分数。这些是论文成绩。不能将摘要中的百分比表述直接换成所有模型至少提高若干个百分点，也不能自行把分类分数等权平均代替原表 Avg。[1，PDF 第 6 页 Table 1]"),
        t(["Qwen 消融设置", "Avg"], [["完整 MetaMem", "71.90"], ["去掉自反思", "70.02"], ["去掉演化", "68.33"]], [12.5, 4.0]),
        p("表 2  原论文 Table 2 的 Qwen 消融结果，支持反思和演化在该设置下的作用。总体提升不保证各子任务都提升，例如 Llama 单会话偏好从 70.24 降至 66.12。离线构建成本也不能被推理阶段耗时表覆盖。[1，PDF 第 6、7、11 页]"),
    ]),
    ("4 本地简化试点与真实规则学习", [
        p("仓库试点使用 BM25 检索原始消息块，保留时间与角色，固定 top-k 20、事实正文最多 12000 字符。它没有实现 LightMem 的压缩、主题聚合或长期维护，也没有复现原论文向量检索配置。模型为 Qwen2.5-7B-Instruct Q4_K_M，原机器服务上下文为 8192。[3，4]"),
        p("固定 cleaned S 来源版本与 SHA-256 后，从 500 条中选择 20 训练、10 开发、20 保留测试。所选三集合的问题与会话 ID、会话内容、完整历史没有直接重叠；源数据没有稳定 user_id，不能把这一检查称为用户级独立性证明。测试尚未运行。"),
        t(["配置", "事实", "经验"], [["B0", "共享固定检索结果", "显式空规则"], ["B1", "共享固定检索结果", "预先固定的人工规则"], ["B2", "共享固定检索结果", "仅从训练学习并由开发集选择的规则"]], [2.0, 7.0, 7.5]),
        p("真实学习采用 20 题、每题 2 次、batch 5、1 epoch。唯一混合正误题组是 Apex 旧目标题，触发 2 次反思、1 次提案与 1 次合并。保存的一条规则为："),
        p("Prioritize recent memories, cross-validate facts, and understand context and chronology unless historical context is required."),
        p("这条规则强调近期、交叉验证与历史语境。它来自 v1 训练；后续 v2 拒答政策只重评旧回答，不构成重新训练。原验证因为存在 uncertain，没有完整有效覆盖候选，未正式选定最终 B2。不能把一次规则生成直接称为效果已证实。[4]"),
        t(["第一阶段稳定性", "Yes", "No", "uncertain"], [["B0 空规则", "17", "13", "0"], ["冻结学习规则", "18", "11", "1"]], [8.0, 2.8, 2.8, 2.9]),
        p("每组 30 条重复观察来自同一批 10 道开发题，各重复 3 次。上述为保守复核结果，不能将 uncertain 当错，也不能声称有 30 道独立题；明信片的旧优势没有稳定重现。[5]"),
    ]),
    ("5 第二阶段的检索诊断", [
        p("第一阶段回接原历史发现，晚餐食材和首饰戒指确有漏检，但明信片累计数量证据已经在输入中。因此第二阶段分别做参考辅助的真实历史补证据诊断，以及不读答案的通用角色预算检索对照。[6]"),
        t(["补证据案例", "原检索判对", "补齐后判对", "解释"], [["晚餐", "1/3", "3/3", "真正食材改善了内容对齐"], ["首饰", "0/3", "1/3", "匹配数字的回答仍漏戒指，依据不完整"]], [2.5, 3.0, 3.0, 8.0]),
        p("这两题的 oracle 使用参考辅助定位真实历史，且改变证据呈现与长度，只用于诊断，不是可部署的自动检索。数字匹配不能替代正确的计数依据，也不能从两题推定全局瓶颈。"),
        p("唯一自动候选沿 BM25 排序给 user 与 other 各 10 块、6000 字符，不转移剩余额度。它没有补回已知缺失食材或戒指，事实字符数和其他上下文也发生变化，因此不能把分数差异只归因于角色。"),
        t(["组别", "检索", "规则", "判对次数"], [["A", "原 BM25", "无", "16/30"], ["B", "原 BM25", "旧学习规则", "17/30"], ["C", "角色预算", "无", "15/30"], ["D", "角色预算", "旧学习规则", "17/30"]], [1.5, 5.5, 5.0, 4.5]),
        p("四组均完整有效覆盖，都是十题重复三次。C 相对 A 仅目标对 1 次、仅基线对 2 次；D 相对 B 各 1 次。角色预算没有一致收益，保留为可选负结果，没有替换默认检索。全阶段 132 个回答与 2 次预热共 134 次本地请求。[6]"),
    ]),
    ("6 第三阶段的回答约束与代价", [
        p("第三阶段固定原 BM25 事实和量化 7B，两组均无规则，仅比较原提示词与人工设计的证据约束候选。候选要求分清对象、事件、时间与计数，并在证据不足时停止猜测。十题各三次，共 60 个新回答，比较使用本轮重新生成的基线。[7]"),
        t(["配置", "第 1 轮", "第 2 轮", "第 3 轮", "Yes 合计", "输出 token"], [["原提示词", "6", "5", "6", "17/30", "9712"], ["候选", "5", "5", "5", "15/30", "3116"]], [4.1, 2.1, 2.1, 2.1, 2.7, 3.4]),
        p("两组有效覆盖均 30/30；配对同对 14、同错 12、仅候选对 1、仅原对 3。候选改善一次明信片累计计算，但晚餐三轮均退步，海报与 Sapiens 的六次拒答机会仍全部失败。输出减少约 67.9%，这是本次工程观测，不等于质量全面提高。"),
        t(["独立依据质量", "原提示词", "候选"], [["事实支持 支持/不支持/U", "9/19/2", "13/12/5"], ["引用存在 有效/无效/无引用", "8/0/22", "28/2/0"], ["引用支持 支持/不支持/U/NA", "4/3/1/22", "12/15/3/0"]], [9.0, 3.75, 3.75]),
        p("表 3  每组均以 30 个回答为范围。U 表示不确定，NA 表示不适用，不能将二者默认算错。28 个回答中的引用 ID 存在，不表示相应事实支持结论；例如海报题用真实的研究会议片段推断本科课程海报地点，仍是跨事件错误。"),
        p("主评后对 38 个回答作一次独立复核，正确性标签一致，但质量有 23 个维度分歧，涉及 10 个回答；未解决的分歧保留 uncertain。评审仍来自 AI，精确部署版本未知，无人工金标准。最终保留候选与负结果，不切换默认提示词。[7]"),
    ]),
    ("7 局限与下一步", [
        p("论文与个人试点的设置差异包括事实构建、模型规模、量化、judge、训练规模和采样条件。当前试点支持对具体机制的学习与诊断，不能用分数差异证明复现或否定原论文。十道开发题已经反复参与方法设计，后续不能把它们改称独立泛化测试。"),
        p("现有证据同时支持检索遗漏、证据利用和评判口径问题。衣物参考数量的对象范围、观鸟问题的时间异常、AI 质量复核分歧仍需集中核对。应先保存人的判断及理由，再固定候选、数据清单、模型、生成预算与正确性/质量政策，随后决定是否开展一次独立测试。"),
        p("讲解自检  能否说明 M 与 E 的区别；为什么全错题组可能不更新；为什么原验证没有选定 B2；为什么真实引用与正确数字仍可能对应错误依据；为什么十题三次不等于三十道独立题。"),
        p("课程衔接  本材料可作为个人论文阅读和已有研究交流的内容。六篇论文的全组整合、本人阅读验收、最终视频及老师后续提交要求仍需分别完成。当前老师通知没有规定必须录制 60 分钟，也没有规定必须完成 500 条五折实验。[8]"),
        p("参考资料"),
        p("[1] Xin et al. MetaMem: Evolving Meta-Memory for Knowledge Utilization through Self-Reflective Symbolic Optimization. Findings of ACL 2026, 5473–5492. " + PAPER),
        p("[2] OpenBMB/MetaMem，固定提交 66232c1cb33d3d7939d323b431f21c741785142e。https://github.com/OpenBMB/MetaMem"),
        p("[3] docs/metamem-local-pilot.md；docs/metamem-code-review.md。"),
        p("[4] docs/metamem-external-ai-run-20261015.md；docs/metamem-strict-abstention-v2-analysis.md。"),
        p("[5] docs/metamem-stage1.md。[6] docs/metamem-stage2.md。[7] docs/metamem-stage3.md 及 reports/metamem-stage3-results.json。"),
        p("[8] 老师原始通知 NLP初步课设与平时作业安排 贺瑞芳 20260920，第 1、4、8–9 页。仓库阶段记录截至 5afb964，本机迁移与模拟校验不增加真实模型成绩。"),
    ]),
]


def slide(title, body, source, notes, **extra):
    return {"title": title, "body": body, "source": source, "notes": notes, **extra}


SLIDES = [
    slide("MetaMem 论文阅读与个人试点", ["长上下文场景结构化记忆建模", "天津大学 NLP 第 8 队  专题 4 论文组 2", "汇报人 __________   2026 年 10 月 19 日"], "MetaMem，Findings of ACL 2026；仓库三阶段记录", "本次先解释论文如何通过自然语言经验改善记忆使用，再介绍仓库中杨垚负责的个人试点。汇报署名需本人确认，已有试点贡献保持原归属。这里讨论的是个人阶段材料，尚无正式测试成绩。", kind="cover"),
    slide("正确片段仍可能形成错误答案", ["三次旅行合计 1800 英里，另一次旅行总计 1200 英里", "另一次旅行的首日 300 英里已经包含在总量中", "正确总量为 3000 英里，重复相加会得到 3300 英里"], "原论文 PDF 第 1 页 Figure 1", "这个例子说明为什么只检索更多事实不够。模型必须保留事件包含关系。元记忆尝试保存避免重复计数的通用经验，具体里程仍属于事实记忆。"),
    slide("事实记忆 M 与元记忆 E", ["M 保存具体历史事实及其背景", "E 保存跨问题可使用的事实处理策略", "回答同时接收 q、M 与冻结的 E，模型权重保持不变"], "原论文 PDF 第 3 页，式 1 至 3", "M 与 E 不是两个名字不同的事实库。1800和1200是具体内容，判断总量与部分关系是使用方式。没有梯度更新不等于没有学习成本，因为仍要使用带答案的数据和多次调用。", table={"headers": ["对象", "示例"], "rows": [["事实 M", "旅行总里程及事件背景"], ["经验 E", "先判定包含关系，避免重复计数"]]}),
    slide("训练反馈与测试回答的边界", ["训练题提供标准答案，支持判分、反思与经验编辑", "开发集选择规则版本，选择过程不继续修改规则", "测试只使用固定问题、检索事实和经验，参考答案单独评分"], "原论文 PDF 第 3、11 页；本地入口接口说明", "标准答案可以在训练反思中使用，但不能进入测试回答提示词。验证用来选规则，测试必须与开发分开。我们的十题已反复用于诊断，所以只能继续称开发集。"),
    slide("MetaMem 的学习循环", [], "原论文 PDF 第 4 页 Figure 2，式 4 至 10", "先从相同事实和当前规则采样多个回答，judge与答案比较，再反思每次成功或失败。随后提出增加、删除或修改规则的动作，同批动作经过冲突过滤后更新元记忆。整个循环更新的是显式经验。", kind="figure"),
    slide("原论文的评估设置", ["LongMemEval 共 500 题，五折每折 350 训练、50 验证、100 测试", "LightMem 构建事实，MiniLM 检索 top-k 20", "Qwen3 30B A3B 与 Llama3.1 70B 回答，Qwen3 235B A22B 判分", "学习 5 轮、batch 50、每题采样 5 次"], "原论文 PDF 第 5、11 页，§4、附录 A.3 和 A.5", "正文400道构建题在附录进一步分成350训练与50验证。原论文使用的模型、事实构建和judge规模都不同于本地试点，不能只看数字大小把两者直接排名。"),
    slide("原论文的平均指标", ["Qwen 绝对提高 4.40 个百分点", "Llama 绝对提高 2.91 个百分点", "总体提高并不保证所有任务均提高"], "原论文 PDF 第 6 页 Table 1，Avg 原值", "图中是原论文成绩而不是小组成绩。Qwen从67.50到71.90，Llama从66.17到69.08。不要把相对百分比与百分点混淆，也不要把分类分数重新等权平均替代原表Avg。", chart={"categories": ["Qwen3 30B A3B", "Llama3.1 70B"], "series": [{"name": "LightMem", "values": [67.50, 66.17]}, {"name": "MetaMem", "values": [71.90, 69.08]}], "max": 100, "unit": "论文 Avg  百分数"}),
    slide("反思和演化的消融证据", ["完整方法的 Avg 为 71.90", "移除反思后为 70.02，移除演化后为 68.33", "这些结果支持对应组件在论文设置下的作用"], "原论文 PDF 第 7 页 Table 2，Qwen 部分", "消融是在相同论文条件下比较方法组件。它不能保证量化小模型同样获益。规则可能泛化过度，构建规则还需要反思和编辑调用，所以不能把方法理解为零成本提示词。", table={"headers": ["设置", "Avg"], "rows": [["完整 MetaMem", "71.90"], ["无自反思", "70.02"], ["无演化", "68.33"]]}),
    slide("本地试点的实现范围", ["原始消息分块与 BM25，固定 top-k 20、12000 字符事实预算", "固定 20 训练、10 开发、20 保留测试，测试尚未运行", "Qwen2.5 7B Q4_K_M，8192 上下文，外部 AI 文件评分", "共享事实比较空规则、固定人工规则与学习规则"], "docs/metamem-local-pilot.md；metamem-local-run-20261015.md", "这是BM25 RAG简化试点，没有LightMem压缩和主题维护。三个集合没有直接共享会话，但源数据无稳定user_id，不能称用户级隔离。后续诊断都限定在十道开发题。"),
    slide("真实学习只形成一条经验", ["20 题每题两次采样，batch 5，1 epoch", "唯一混合正误 Apex 题组触发 2 次反思、1 次提案、1 次合并", "经验强调近期事实、交叉验证与历史语境", "原验证存在 uncertain，未正式选定最终 B2"], "外部 AI v1 实验记录；严格拒答 v2 敏感性分析", "官方实现仅对同题有对有错的组更新，全对和全错都跳过。元数据有20个题号并不表示20题都贡献规则。后续v2重判没有产生新的反思或规则，因此不能称v2重新训练。"),
    slide("规则收益的稳定性检查", ["同一批 10 道开发题，各配置重复 3 次", "旧明信片优势未稳定重现", "不确定项保留，不能当成答错"], "docs/metamem-stage1.md，保守复核统计", "30条是重复观察，不是30道独立题。B2有一次评分分歧，保守统计保留uncertain。微小差距不足以支持稳定泛化收益，仍需查看逐题配对和评分可靠性。", table={"headers": ["配置", "Yes", "No", "uncertain"], "rows": [["B0 空规则", "17", "13", "0"], ["冻结学习规则", "18", "11", "1"]]}),
    slide("自动角色预算未显示稳定收益", ["原检索与角色预算分别比较有无旧规则", "角色预算未补齐已知食材和戒指缺口", "两题补证据仅用于诊断，不能算自动检索提升"], "docs/metamem-stage2.md，主实验与 oracle 分开统计", "四组都在十题上各重复三次。晚餐补证据从1/3到3/3，说明覆盖作用；首饰唯一匹配数字的回答仍漏戒指，所以不能只报成功。自动角色预算没有一致方向，保留为可选负结果。", table={"headers": ["组", "检索", "规则", "Yes"], "rows": [["A", "原 BM25", "无", "16/30"], ["B", "原 BM25", "旧规则", "17/30"], ["C", "角色预算", "无", "15/30"], ["D", "角色预算", "旧规则", "17/30"]]}),
    slide("证据约束提示词的正确性对照", ["事实相同，两组均无规则", "原提示词合计 17/30，候选 15/30", "仅候选对 1 次，仅原提示词对 3 次"], "docs/metamem-stage3.md；reports/metamem-stage3-results.json", "候选要求区分事件、时间、计数并停止猜测。基线本轮重新生成。候选改善明信片第2轮，但晚餐三轮退步。海报和Sapiens仍全部拒答失败，因此没有切换默认提示词。", chart={"categories": ["第 1 轮", "第 2 轮", "第 3 轮"], "series": [{"name": "原提示词", "values": [6, 5, 6]}, {"name": "候选", "values": [5, 5, 5]}], "max": 10, "unit": "每轮十题中的判对次数"}),
    slide("回答更短仍需要检查依据", ["候选输出 token 合计减少约 67.9%", "事实支持为 9/30 与 13/30，不确定项为 2 与 5", "候选 28 条引用 ID 存在，引用支持仅 12 条"], "第三阶段结果，正确性与依据质量分别评审", "图展示实际输出量，不是独立速度基准。引用ID存在只能说明格式和来源可以匹配，不能证明内容支持结论。候选还有15条引用不支持、3条不确定，质量维度不能被压成单一正确率。", chart={"categories": ["原提示词", "候选"], "series": [{"name": "输出 token", "values": [9712, 3116]}], "max": 11000, "unit": "两组各 30 个回答的输出 token 合计"}),
    slide("三个案例说明不同错误来源", ["明信片已有累计 25 的证据，模型仍可能误做减法", "海报引用真实研究会议片段，却混淆为本科课程海报事件", "首饰存在漏检，补齐后仍可能误读日期或计数单位"], "第一至三阶段逐题证据与错误分析", "这三类现象分别涉及证据利用、跨事件推断和检索与利用交织。不能把全部问题归因检索，也不能用最后数字匹配来证明推理成立。晚餐则说明真正食材没有检索到时，规则很难凭空补齐事实。"),
    slide("结论的适用范围", ["十道开发题已反复用于方法设计，独立测试尚未运行", "AI 主评与复核均无人工金标准，精确部署版本未知", "衣物参考口径与观鸟时间异常仍需核对", "角色预算和人工提示词都保留为候选，没有替换默认"], "统一阶段总报告；课程交付差距", "原论文和本地试点不同设置不能混算。三次重复不扩大独立题目数。负结果仍然回答了当前候选是否值得采用，但不能证明所有元记忆、所有检索或提示词方法均无效。"),
    slide("下一步与讲解自检", ["先集中核对代表性证据、参考口径和 AI 评分分歧", "在看测试内容前冻结候选、配置、预算与一次性评估政策", "本人应解释 M/E、更新筛选、开发与测试边界及负结果", "小组还需六篇论文整合与最终视频，时长以老师通知为准"], "docs/metamem-course-readiness.md；老师原始通知第 1、4 页", "我们已有可解释的研究链条，下一步优先处理人的验收和独立评估边界，不追加第四种方案来追分。课设现有通知要求系统Word、PPT和录好视频，没有明确60分钟。个人材料不代替全组六篇整合。"),
    slide("主要参考资料", ["MetaMem，Findings of ACL 2026，5473–5492", "OpenBMB/MetaMem 固定版本 66232c1", "仓库第一至三阶段报告与对应结构化结果", "老师原始通知第 1、4、8–9 页"], PAPER, "论文完整链接为https://aclanthology.org/2026.findings-acl.270/。方法见PDF第3至5页，主表第6页，消融第7页，五折设置第11页。代码来自OpenBMB/MetaMem固定提交66232c1cb33d3d7939d323b431f21c741785142e。已有个人试点负责人杨垚；汇报人应在阅读核验后填写署名。"),
]


def font(run, size=12, bold=False, east="SimSun"):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._element.get_or_add_rPr()
    f = props.rFonts
    if f is None:
        f = OxmlElement("w:rFonts"); props.insert(0, f)
    f.set(qn("w:eastAsia"), east)


def add_table(doc, block):
    rows = [block["headers"], *block["rows"]]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block["widths"] or [16.5 / len(rows[0])] * len(rows[0])
    for ci, width in enumerate(widths):
        table.columns[ci].width = Cm(width)
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), "DCE6F1" if ri == 0 else ("F4F7FA" if ri % 2 == 0 else "FFFFFF")); tcpr.append(fill)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge); e.set(qn("w:w"), "90"); e.set(qn("w:type"), "dxa"); margins.append(e)
            tcpr.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.12
            if ci > 0 and len(str(value)) < 20:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(para.add_run(str(value)), 10.5, ri == 0)
        trpr = table.rows[ri]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); trpr.append(cant)
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "D9D9D9"); borders.append(e)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    BUILD.mkdir(parents=True, exist_ok=True)
    fig = BUILD / "metamem-figure2.png"
    if not fig.exists():
        exe = shutil.which("pdftoppm")
        if not exe:
            raise RuntimeError("Use the Codex runtime Poppler on PATH to extract the cited figure")
        subprocess.run([exe, "-f", "4", "-singlefile", "-scale-to", "1800", "-png", str(ROOT / "papers/structured-memory-group2/05_MetaMem.pdf"), str(BUILD / "paper-method-page")], check=True)
        Image.open(BUILD / "paper-method-page.png").crop((150, 150, 1120, 741)).save(fig)
    evidence = json.loads((ROOT / "reports/metamem-stage3-results.json").read_text(encoding="utf-8"))
    assert evidence["overall"]["baseline"]["Yes"] == 17
    assert evidence["overall"]["candidate"]["Yes"] == 15
    assert evidence["independent_main_questions"] == 10 and evidence["human_acceptance"] is None
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.05)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 23), ("Heading 1", 17), ("Heading 2", 14)):
        style = doc.styles[name]; style.font.name = "Microsoft YaHei"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(12)
    for style in doc.styles:
        for border in list(style._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    markdown = []
    for page_no, (heading, blocks) in enumerate(PAGES):
        if page_no:
            doc.add_page_break()
        title = doc.add_paragraph(style="Title" if page_no == 0 else "Heading 1")
        if page_no == 0:
            borders = OxmlElement("w:pBdr")
            for edge in ("top", "bottom", "left", "right", "between"):
                border = OxmlElement("w:" + edge); border.set(qn("w:val"), "nil"); borders.append(border)
            title._p.get_or_add_pPr().append(borders)
        font(title.add_run(heading), 23 if page_no == 0 else 17, True, "Microsoft YaHei")
        markdown.append(("# " if page_no == 0 else "## ") + heading + "\n")
        for block in blocks:
            if block["type"] == "paragraph":
                para = doc.add_paragraph()
                font(para.add_run(block["text"]))
                markdown.append(block["text"] + "\n")
            elif block["type"] == "table":
                add_table(doc, block)
                markdown.append("| " + " | ".join(block["headers"]) + " |\n|" + "---|" * len(block["headers"]))
                markdown.extend("| " + " | ".join(row) + " |" for row in block["rows"])
                markdown.append("")
            else:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.add_run().add_picture(str(fig), width=Cm(16.3))
                caption = doc.add_paragraph(); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                font(caption.add_run(block["caption"]), 9.5)
                markdown.append(block["caption"] + "\n")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.core_properties.title = PAGES[0][0]
    doc.core_properties.author = ""
    doc.core_properties.subject = "MetaMem 个人论文阅读和已有三阶段试点整理"
    doc.save(ROOT / "reports" / (STEM + ".docx"))
    (ROOT / "reports" / (STEM + ".md")).write_text("\n".join(markdown), encoding="utf-8", newline="\n")
    slide_content = {"title": PAGES[0][0], "date": "2026-10-19", "slides": SLIDES,
                     "figure": "tmp/metamem-materials-20261019/metamem-figure2.png", "font": "Microsoft YaHei"}
    (ROOT / "reports/metamem-course-slides.json").write_text(json.dumps(slide_content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({"document": STEM + ".docx", "planned_sections": len(PAGES), "slides": len(SLIDES)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
